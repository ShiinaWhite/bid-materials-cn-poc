"""Read-only, streaming extraction into disposable staging. No SQLite writes.

Optional parsers are imported only inside their format handler. OCR uses the
existing Umi engine, one necessary page at a time, in memory (no page images).

v0.8: DOCX embedded images become first-class content. Anchors are built in
one document walk shared with the native-text units (order, heading breadcrumb
and bounded surrounding text), then every raster image is OCR'd exactly once:
sequential (concurrency 1), resumable by cache key (source SHA + image SHA +
engine + params), with per-image failure isolation — one bad image never
fails the document.
"""
import base64
from contextlib import closing
import json
from pathlib import Path
import struct
import time
import zipfile

from .common import atomic_write, digest, normalize
from .ocr import Ocr

SUPPORTED = {'pdf', 'docx', 'xlsx', 'txt', 'md'}
MAX_TEXT = 100_000_000
MAX_UNITS = 50_000
# Embedded-image budgets: bounded work on modest hardware, one image at a time.
MAX_IMAGES = 2_000
IMAGE_MAX_BYTES = 32 * 1024 * 1024
IMAGE_SKIP_MIN_BYTES = 2 * 1024  # tiny blobs: inline icons/rules, not material
IMAGE_SKIP_MIN_DIM = 40          # pixels; smaller than this is decorative
OCR_ENGINE = 'umi-ocr'
OCR_PARAMS_VERSION = 1


class Unavailable(RuntimeError):
    pass


def check_office(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 20_000 or sum(e.file_size for e in entries) > 256 * 1024 * 1024:
            raise ValueError('Office container exceeds safe extraction budget')
        if any(e.filename.lower().endswith('vbaproject.bin') for e in entries):
            raise ValueError('macro-bearing Office container is not supported')


class PageOcr:
    def __init__(self, config):
        self.c, self.engine = config, None

    def recognize(self, data):
        if self.engine is None:
            self.c.logs_dir.mkdir(parents=True, exist_ok=True)
            engine = Ocr(self.c)
            try:
                engine.start()
            except Exception:
                engine.stop()
                raise
            self.engine = engine
        try:
            response = self.engine.request('/api/ocr', dict(
                base64=base64.b64encode(data).decode('ascii'), options={
                    'ocr.language': '简体中文', 'ocr.maxSideLen': 2000,
                    'tbpu.parser': 'multi_para', 'data.format': 'dict'}), timeout=120)
            if not isinstance(response, dict) or response.get('code') not in (100, 101):
                raise RuntimeError('Umi page OCR failed')
            blocks = response.get('data', []) if response['code'] == 100 else []
            if not isinstance(blocks, list) or any(not isinstance(b, dict) or not isinstance(b.get('text'), str) for b in blocks):
                raise RuntimeError('Umi returned malformed text blocks')
            return '\n'.join(b['text'] for b in blocks)
        except OSError:
            self.close()
            raise

    def close(self):
        if self.engine is not None:
            engine, self.engine = self.engine, None
            engine.stop()


def native_sufficient(text):
    visible = ''.join(c for c in text if not c.isspace())
    return len(visible) >= 30 and visible.count('\ufffd') <= len(visible) * .02


def image_mime(data):
    if data[:8] == b'\x89PNG\r\n\x1a\n':
        return 'image/png'
    if data[:3] == b'\xff\xd8\xff':
        return 'image/jpeg'
    if data[:6] in (b'GIF87a', b'GIF89a'):
        return 'image/gif'
    if data[:2] == b'BM':
        return 'image/bmp'
    if data[:4] == b'RIFF' and data[8:12] == b'WEBP':
        return 'image/webp'
    return None


def image_dimensions(data):
    """Header-only width/height for the raster formats Umi accepts."""
    try:
        if data[:8] == b'\x89PNG\r\n\x1a\n' and data[12:16] == b'IHDR':
            return struct.unpack('>II', data[16:24])
        if data[:6] in (b'GIF87a', b'GIF89a'):
            return struct.unpack('<HH', data[6:10])
        if data[:2] == b'BM':
            return struct.unpack('<ii', data[18:26])
        if data[:3] == b'\xff\xd8\xff':  # first SOFn marker carries the frame size
            offset = 2
            while offset + 9 < len(data):
                if data[offset] != 0xFF:
                    offset += 1
                    continue
                marker, size = data[offset + 1], struct.unpack('>H', data[offset + 2:offset + 4])[0]
                if marker in (0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF):
                    height, width = struct.unpack('>HH', data[offset + 5:offset + 9])
                    return width, height
                offset += 2 + size
    except (IndexError, struct.error):
        pass
    return None


def image_cache_key(source_sha256, image_sha256):
    return f'{source_sha256}|{image_sha256}|{OCR_ENGINE}|v{OCR_PARAMS_VERSION}'


_A_NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
_R_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
_V_NS = '{urn:schemas-microsoft-com:vml}'


def _element_image_refs(element, part):
    """DrawingML blips plus legacy VML imagedata, resolved to zip media paths."""
    refs = []
    for blip in element.findall('.//' + _A_NS + 'blip'):
        rid = blip.get(_R_NS + 'embed') or blip.get(_R_NS + 'link')
        if rid:
            refs.append(rid)
    for imagedata in element.findall('.//' + _V_NS + 'imagedata'):
        rid = imagedata.get(_R_NS + 'id')
        if rid:
            refs.append(rid)
    resolved = []
    for rid in refs:
        try:
            rel = part.rels[rid]
            if rel.is_external:
                continue
            # partname is absolute inside the package ('/word/media/x.png').
            resolved.append((str(rel.target_part.partname).lstrip('/'), rid))
        except KeyError:
            continue
    return resolved


def _docx_events(document):
    """One ordered walk: heading / text-line / image events. The native unit
    builder and the image anchor builder share this single source of order."""
    from docx.text.paragraph import Paragraph
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            level = heading_level(block)
            if level is not None and level <= 9:
                yield dict(kind='heading', level=level, text=block.text)
            for media, rid in _element_image_refs(block._p, document.part):
                yield dict(kind='image', media=media, rid=rid)
            if block.text.strip():
                yield dict(kind='text', text=block.text)
        else:
            for row in block.rows:
                line = '\t'.join(c.text for c in row.cells)
                if line.strip():
                    yield dict(kind='text', text=line)
            for media, rid in _element_image_refs(block._tbl, document.part):
                yield dict(kind='image', media=media, rid=rid)


def pdf_units(path, progress, ocr):
    try:
        import pymupdf
    except ImportError as error:
        raise Unavailable('PDF extraction unavailable: install optional PyMuPDF locally') from error
    with pymupdf.open(path) as document:
        if document.is_encrypted:
            raise ValueError('encrypted PDF is unsupported without a separately supplied decrypted original')
        total = len(document)
        progress(0, total, 'page')
        for i, page in enumerate(document):
            text, method, error = '', 'failed', None
            try:
                text = page.get_text('text', sort=True)
                if native_sufficient(text):
                    method = 'native'
                else:
                    # At most 2000 px on the long edge; discard the pixmap after
                    # this page. Retain scant native text if OCR fails.
                    zoom = min(2.5, 2000 / max(page.rect.width, page.rect.height))
                    scanned = ocr.recognize(page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).tobytes('png'))
                    text = scanned if scanned.strip() else text
                    method = 'ocr' if text.strip() else 'empty'
            except Exception as failure:
                error = str(failure)[:1000]
            yield dict(unit_type='page', page_no=i + 1, text_raw=text, page_status=method, error=error)
            progress(i + 1, total, 'page')


def heading_level(paragraph):
    # python-docx safely handles XML/relationships. Follow basedOn styles too,
    # including localized/custom headings that carry outlineLvl.
    from docx.oxml.ns import qn
    props = paragraph._p.pPr
    outline = props.find(qn('w:outlineLvl')) if props is not None else None
    if outline is not None:
        return int(outline.get(qn('w:val'))) + 1
    style, seen = paragraph.style, set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        if style.name and style.name.lower().startswith('heading '):
            try:
                return int(style.name.split()[-1])
            except ValueError:
                pass
        props = style.element.pPr
        outline = props.find(qn('w:outlineLvl')) if props is not None else None
        if outline is not None:
            return int(outline.get(qn('w:val'))) + 1
        style = style.base_style
    return None


def docx_units(path, progress, warnings, images=None, ocr=None, resume=None, source_sha256=''):
    """Native text units, then (when an images list is supplied) the embedded
    image phase: anchors, per-image OCR with resume, and image text units.
    Without an images list this stays the v0.6 native-text-only generator that
    project intake text reuse expects."""
    import hashlib
    try:
        from docx import Document
        from docx.text.paragraph import Paragraph  # noqa: F401
    except ImportError as error:
        raise Unavailable('DOCX extraction unavailable: install python-docx') from error
    check_office(path)
    document = Document(path)
    events = list(_docx_events(document))

    headings, group, length = {}, [], 0

    def section():
        return ' → '.join(headings[k] for k in sorted(headings)) or '正文'

    for event in events:
        if event['kind'] == 'heading':
            if group:
                yield dict(unit_type='paragraph_group', section_path=section(),
                           text_raw='\n'.join(group), page_status='native')
                group, length = [], 0
            headings = {k: v for k, v in headings.items() if k < event['level']}
            headings[event['level']] = event['text']
        elif event['kind'] == 'text':
            line = event['text']
            if group and length + len(line) > 6000:
                yield dict(unit_type='paragraph_group', section_path=section(),
                           text_raw='\n'.join(group), page_status='native')
                group, length = [], 0
            group.append(line)
            length += len(line)
    if group:
        yield dict(unit_type='paragraph_group', section_path=section(),
                   text_raw='\n'.join(group), page_status='native')

    if images is None:
        return
    # ---- embedded image anchors: replay the walk to bind each image to its
    # heading breadcrumb and the nearest surrounding text.
    anchors, heading_state, recent = [], {}, []
    for i, event in enumerate(events):
        if event['kind'] == 'heading':
            heading_state = {k: v for k, v in heading_state.items() if k < event['level']}
            heading_state[event['level']] = event['text']
        elif event['kind'] == 'text':
            recent.append(event['text'].strip())
        else:
            following = [e['text'].strip() for e in events[i + 1:] if e['kind'] == 'text'][:2]
            anchors.append(dict(
                media_path=event['media'], relationship_id=event['rid'],
                document_order=len(anchors) + 1,
                section_path=' → '.join(heading_state[k] for k in sorted(heading_state)) or '正文',
                nearest_heading=heading_state[max(heading_state)] if heading_state else None,
                preceding_text=' / '.join([t for t in recent[-2:] if t]) or None,
                following_text=' / '.join([t for t in following if t]) or None))
    if len(anchors) > MAX_IMAGES:
        raise ValueError('document exceeds safe embedded image extraction budget')
    if not anchors:
        return
    progress(0, len(anchors), 'image')
    done_by_sha = {}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for anchor in anchors:
            record = dict(anchor, mime=None, byte_size=0, sha256=None, width_px=None, height_px=None,
                          ocr_status='pending', ocr_text_raw=None, ocr_text_norm=None,
                          ocr_error=None, ocr_elapsed_ms=None, ocr_cache_key=None, reused=False)
            images.append(record)
            try:
                media = anchor['media_path']
                if media not in names:
                    record.update(ocr_status='skipped', ocr_error='media entry missing from container')
                    continue
                info = archive.getinfo(media)
                record['byte_size'] = info.file_size
                if media.lower().endswith(('.wmf', '.emf')):
                    record.update(ocr_status='skipped', ocr_error='vector metafile is not OCR material')
                    continue
                if info.file_size > IMAGE_MAX_BYTES:
                    record.update(ocr_status='skipped', ocr_error='image exceeds per-image size budget')
                    continue
                data = archive.read(media)  # exactly one image resident at a time
                try:
                    record['sha256'] = hashlib.sha256(data).hexdigest()
                    record['mime'] = image_mime(data)
                    dimensions = image_dimensions(data)
                    if dimensions:
                        record['width_px'], record['height_px'] = dimensions
                    key = image_cache_key(source_sha256, record['sha256'])
                    record['ocr_cache_key'] = key
                    prior = (resume or {}).get(key)
                    if prior is not None:
                        record.update(ocr_status='done', ocr_text_raw=prior.get('text_raw') or '',
                                      ocr_text_norm=prior.get('text_norm') or '', reused=True,
                                      ocr_elapsed_ms=prior.get('elapsed_ms'))
                        done_by_sha[record['sha256']] = record
                        continue
                    if record['sha256'] in done_by_sha:  # same picture reused later in the document
                        twin = done_by_sha[record['sha256']]
                        record.update(ocr_status=twin['ocr_status'], ocr_text_raw=twin['ocr_text_raw'],
                                      ocr_text_norm=twin['ocr_text_norm'], ocr_error=twin['ocr_error'],
                                      ocr_elapsed_ms=twin['ocr_elapsed_ms'], reused=True)
                        continue
                    if info.file_size < IMAGE_SKIP_MIN_BYTES or (
                            dimensions and min(dimensions) < IMAGE_SKIP_MIN_DIM):
                        record.update(ocr_status='skipped', ocr_error='below decorative-size threshold')
                        continue
                    if record['mime'] is None:
                        record.update(ocr_status='skipped', ocr_error='unrecognized raster image format')
                        continue
                    started = time.monotonic()
                    text = ocr.recognize(data)
                    record.update(ocr_status='done', ocr_text_raw=text,
                                  ocr_text_norm=normalize(text) if text else '',
                                  ocr_elapsed_ms=int((time.monotonic() - started) * 1000))
                    done_by_sha[record['sha256']] = record
                finally:
                    del data
            except Exception as failure:  # one image never fails the document
                record.update(ocr_status='failed', ocr_error=str(failure)[:500])
            progress(record['document_order'], len(anchors), 'image')
    for record in images:
        if record['ocr_status'] == 'done' and (record['ocr_text_raw'] or '').strip():
            yield dict(unit_type='embedded_image', section_path=record['section_path'],
                       text_raw=record['ocr_text_raw'], page_status='ocr',
                       _image_ref=(record['document_order'], record['media_path']))


def xlsx_units(path, progress):
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise Unavailable('XLSX extraction unavailable: install openpyxl') from error
    check_office(path)
    # No links, VBA, Excel runtime, DDE or recalculation. Formula text is inert.
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)) as formulas, \
            closing(openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)) as cached:
        total = len(formulas.worksheets)
        progress(0, total, 'sheet')
        for si, sheet in enumerate(formulas.worksheets):
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 5_000_000:
                raise ValueError('XLSX declared cell region exceeds safe extraction budget')
            rows, start, end, width, length = [], None, None, 1, 0

            def unit():
                return dict(unit_type='sheet_chunk', sheet_name=sheet.title, row_start=start, row_end=end,
                            cell_range=f'A{start}:{get_column_letter(width)}{end}', text_raw='\n'.join(rows), page_status='native')

            for ri, (row, values) in enumerate(zip(sheet.iter_rows(values_only=True), cached[sheet.title].iter_rows(values_only=True)), 1):
                cells = []
                for ci, value in enumerate(row, 1):
                    if value is None:
                        continue
                    text = str(value)
                    if isinstance(value, str) and value.startswith('=') and ci <= len(values) and values[ci-1] is not None:
                        text += ' [cached: ' + str(values[ci-1]) + ']'
                    cells.append(f'{get_column_letter(ci)}{ri}: {text}')
                if not cells:
                    if rows:
                        yield unit()
                    rows, start, end, width, length = [], None, None, 1, 0
                    continue
                line = '\t'.join(cells)
                if rows and (len(rows) >= 40 or length + len(line) > 6000):
                    yield unit()
                    rows, start, end, width, length = [], None, None, 1, 0
                if start is None:
                    start = ri
                end, width = ri, max(width, len(row))
                rows.append(line)
                length += len(line)
            if rows:
                yield unit()
            progress(si + 1, total, 'sheet')


def extract(source, stage, config, progress=lambda *args: None, ocr_factory=PageOcr, resume=None):
    """Stage bounded JSONL. Source is resolved by manager, never from Web input.

    Writes units.jsonl (all searchable units) and, for DOCX, images.jsonl (the
    embedded-image registry records) next to it. The returned summary is a
    structured, displayable account of what did and did not get extracted."""
    path = Path(source['path'])
    if not path.is_file() or path.stat().st_size != source['size_bytes'] or digest(path) != source['sha256']:
        raise ValueError('source size/SHA mismatch; original must be repaired before extraction')
    ext, warnings, images = source['extension'], [], []
    if ext not in SUPPORTED:
        return dict(status='unsupported', error='stored original supported; content extraction deferred',
                    unit_count=0, text_length=0,
                    summary=dict(native_text=dict(units=0, chars=0), embedded_images=None))
    ocr = ocr_factory(config)
    count, length, failed = 0, 0, 0
    page_stats = {'native': 0, 'ocr': 0, 'empty': 0, 'failed': 0}
    ocr_units, ocr_chars = 0, 0
    result = None
    try:
        if ext == 'pdf':
            units = pdf_units(path, progress, ocr)
        elif ext == 'docx':
            units = docx_units(path, progress, warnings, images=images, ocr=ocr,
                               resume=resume, source_sha256=source['sha256'])
        elif ext == 'xlsx':
            units = xlsx_units(path, progress)
        else:
            def text_units():
                with path.open(encoding='utf-8-sig') as stream:
                    group = []
                    for line in stream:
                        group.append(line.rstrip('\n'))
                        if len(group) >= 40:
                            yield dict(unit_type='paragraph_group', section_path='正文', text_raw='\n'.join(group), page_status='native')
                            group = []
                    if group:
                        yield dict(unit_type='paragraph_group', section_path='正文', text_raw='\n'.join(group), page_status='native')
            units = text_units()
        with stage.open('w', encoding='utf-8', newline='\n') as output:
            for item in units:
                count += 1
                length += len(item['text_raw'])
                failed += item.get('page_status') == 'failed'
                if item.get('page_status') in page_stats:
                    page_stats[item['page_status']] += 1
                if item.get('unit_type') == 'embedded_image':
                    ocr_units += 1
                    ocr_chars += len(item['text_raw'])
                if count > MAX_UNITS or length > MAX_TEXT:
                    raise ValueError('document exceeds safe text/unit extraction budget')
                item['text_norm'] = normalize(item['text_raw'])
                # Compute and fsync text derivatives outside the repository
                # lock. Only manager may rename/publish this staging folder.
                atomic_write(stage.parent/'projections'/f'unit-{count:06d}.txt', item['text_norm'])
                output.write(json.dumps(item, ensure_ascii=False) + '\n')
        if digest(path) != source['sha256']:
            raise ValueError('source SHA changed during extraction')
        image_stats = None
        if ext == 'docx':
            with (stage.parent/'images.jsonl').open('w', encoding='utf-8', newline='\n') as output:
                for record in images:
                    output.write(json.dumps(record, ensure_ascii=False) + '\n')
            image_stats = dict(
                total=len(images),
                done=sum(1 for r in images if r['ocr_status'] == 'done'),
                failed=sum(1 for r in images if r['ocr_status'] == 'failed'),
                skipped=sum(1 for r in images if r['ocr_status'] == 'skipped'),
                pending=sum(1 for r in images if r['ocr_status'] == 'pending'),
                reused=sum(1 for r in images if r.get('reused')),
                ocr_chars=ocr_chars, ocr_units=ocr_units)
        problems = list(warnings)
        if failed:
            problems.append(f'{failed} pages failed')
        if image_stats and image_stats['failed']:
            problems.append(f"{image_stats['failed']} embedded image(s) OCR failed")
        summary = dict(native_text=dict(units=count - ocr_units, chars=length - ocr_chars),
                       embedded_images=image_stats, pages=page_stats, problems=problems)
        result = dict(status='partial' if problems else 'complete', unit_count=count, text_length=length,
                      error='; '.join(problems) or None, summary=summary)
    finally:
        try:
            ocr.close()
        except Exception as error:
            config.log('project.document.ocr_cleanup_failed', error=str(error))
            if result is not None:
                result['status'] = 'partial'
                result['error'] = '; '.join(filter(None, [result.get('error'), 'OCR cleanup failed: '+str(error)[:1000]]))
                result['summary'] = dict(result.get('summary') or {}, problems=[*(result.get('summary') or {}).get('problems', []), 'OCR cleanup failed'])
    return result

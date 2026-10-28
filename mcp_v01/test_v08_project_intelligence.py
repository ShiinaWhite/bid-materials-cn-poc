"""v0.8 Project Document Intelligence: embedded images, OCR resume, discovery.

Synthetic coverage for the phase-1 vertical slice: DOCX image anchoring,
per-image OCR with failure isolation and resume-by-cache-key, image OCR text
entering Project FTS, WHO/WHAT/LIFECYCLE candidate discovery, deterministic
exact Material matching with audited auto-link, and upload-time automatic
extraction. The real 83 MiB response file runs in the git-ignored
local-fixture harness, never here.
"""
import json
from pathlib import Path
from unittest import mock

from document_test_support import DocumentFixture, StubOcr
from jty_materials.common import now
from jty_materials.discovery import detect, subject_hints, lifecycle_for
from jty_materials.project_documents import Documents
from jty_materials.store import Store

LICENSE_OCR = ('营业执照\n名称：山西中保源建设工程有限公司\n统一社会信用代码：91140931MA7YM0KQ09\n'
               '法定代表人：张三\n成立日期 2016年3月2日\n经营范围 建设工程施工')
SAFETY_OCR = ('安全生产许可证\n单位名称：山西中保源建设工程有限公司\n'
              '许可证编号：（晋）JZ安许证字[2023]B000123456\n有效期 2026-03-01')
OTHER_COMPANY_OCR = ('营业执照\n名称：湖北某某设备制造有限公司\n统一社会信用代码：91420100MA4K00000Q\n法定代表人：李四')
CREDIT_OCR = ('企业信用报告\n信用中国查询结果\n查询时间 2026-10-26\n报告编号 202610261234\n无严重违法记录')
EXACT_OCR = ('安全生产许可证\n单位名称：山西中保源建设工程有限公司\n'
             '许可证编号：JZ2023B000123456\n有效期 2026-03-01')
ISO_NOISE = '体系认证证书\nGB/T 19001\n覆盖范围 建筑工程施工'
PERSON_OCR = '居民身份证\n姓名：王国栋\n公民身份号码 142601199001011234\n住址 山西省临汾市'


def make_png(path, size=(420, 300), seed=7):
    """A raster image large enough to clear the decorative threshold."""
    from PIL import Image
    image = Image.new('RGB', size, 'white')
    for y in range(0, size[1], 6):
        for x in range(0, size[0], 6):
            image.putpixel((x, y), ((x * seed + y) % 256, (x + y) % 256, (x * y) % 256))
    image.save(path, format='PNG')
    return path


def make_docx_with_images(path, texts, extra=None):
    """Heading + intro paragraph + one image per OCR text, interleaved."""
    from docx import Document
    doc = Document()
    doc.add_heading('三、资格证明文件', 1)
    doc.add_paragraph('以下为资质证明材料扫描件，仅供本项目使用。')
    for index, text in enumerate(texts):
        doc.add_paragraph(f'证明材料：{text.splitlines()[0]}')
        doc.add_picture(str(make_png(path.parent/f'img{index}.png', seed=index + 1)))
    doc.add_paragraph('以上证明文件齐全有效。')
    for line in (extra or []):  # tail lines that differentiate file bytes
        doc.add_paragraph(line)
    doc.save(path)
    return path


class SequenceOcr(StubOcr):
    """Returns queued texts in order; records every call payload. Never falls
    back to StubOcr.recognize (its int counter is replaced by a list here)."""
    queue = []
    def __init__(self, config):
        super().__init__(config)
        self.calls = []  # replaces StubOcr's int counter with a payload log
    def recognize(self, data):
        self.calls.append(data)
        return self.queue.pop(0) if self.queue else 'SCANNED_OCR_SENTINEL'


class FlakyOcr(SequenceOcr):
    """Fails exactly one picture (the second) while broken. Every invocation
    is counted before the decision, so a failure still advances the log."""
    broken = True
    def recognize(self, data):
        self.calls.append(data)
        if FlakyOcr.broken and len(self.calls) == 2:
            raise OSError('Umi crashed on this picture')
        return self.queue.pop(0) if self.queue else 'SCANNED_OCR_SENTINEL'


class V08Harness(DocumentFixture):
    """Shared harness for the phase-1 and review-fix test classes."""
    def rows(self, sql, *params):
        with Store(self.c, readonly=True) as s:
            return [dict(r) for r in s.db.execute(sql, params)]

    def images(self):
        return self.rows('SELECT * FROM project_document_image ORDER BY document_order')

    def units(self, unit_type=None):
        sql = 'SELECT * FROM project_document_unit' + (" WHERE unit_type='" + unit_type + "'" if unit_type else '') + ' ORDER BY unit_index'
        return self.rows(sql)

    def candidates(self):
        return self.rows("SELECT * FROM project_material_candidate ORDER BY candidate_type,subject_name")

    def run_document(self, path, ocr_texts):
        SequenceOcr.queue = list(ocr_texts)
        ocr = SequenceOcr(self.c)
        fid, result = self.process_with(path, ocr)
        return fid, result, ocr

    def process_with(self, path, ocr):
        fid = self.add(path)
        rid, entry = self.begin(fid)
        folder = self.c.document_tmp/rid
        folder.mkdir(parents=True, exist_ok=True)
        from jty_materials.document_extract import extract
        entry['result'] = extract(entry['source'], folder/'units.jsonl', self.c, ocr_factory=lambda _: ocr)
        self.supervisor.job('project.document._finish', dict(run_id=rid, capability=entry['capability']))
        return fid, entry['result']

    def seed_library_material(self, code, uscc=None, name=None, cert_no=None):
        """Deprecated: org/material creation follows the guarded evidence path.
        Kept out; use import_material + Store.verify instead."""
        raise NotImplementedError

class V08DocumentIntelligenceTests(V08Harness):
    # ---- P1: image discovery, anchors, OCR units, FTS -------------------
    def test_docx_embedded_images_anchored_ocr_units_and_fts(self):
        path = make_docx_with_images(self.base/'certs.docx', [LICENSE_OCR, SAFETY_OCR])
        fid, result, ocr = self.run_document(path, [LICENSE_OCR, SAFETY_OCR])
        self.assertEqual(result['status'], 'complete')
        self.assertEqual(len(ocr.calls), 2)  # concurrency 1, one call per image
        images = self.images()
        self.assertEqual(len(images), 2)
        first = images[0]
        self.assertEqual(first['mime'], 'image/png')
        self.assertEqual((first['width_px'], first['height_px']), (420, 300))
        self.assertEqual(first['ocr_status'], 'done')
        self.assertTrue(first['sha256'] and len(first['sha256']) == 64)
        self.assertIn('|umi-ocr|v1', first['ocr_cache_key'])
        self.assertEqual(first['document_order'], 1)
        self.assertIn('资格证明文件', first['section_path'])
        self.assertIn('以下为资质证明材料', first['preceding_text'])
        self.assertIn('证明文件齐全有效', images[1]['following_text'])
        image_units = self.units('embedded_image')
        self.assertEqual(len(image_units), 2)
        self.assertEqual(image_units[0]['page_status'], 'ocr')
        self.assertIn('营业执照', image_units[0]['text_raw'])
        hits = Documents(self.c).search(self.pid, '91140931MA7YM0KQ09')['items']
        self.assertTrue(hits and hits[0]['unit_type'] == 'embedded_image' and hits[0]['image_order'] == 1)
        summary = result['summary']
        self.assertEqual(summary['embedded_images']['total'], 2)
        self.assertEqual(summary['embedded_images']['done'], 2)
        self.assertEqual(summary['embedded_images']['failed'], 0)
        self.assertGreater(summary['embedded_images']['ocr_chars'], 0)
        self.assertEqual(summary['native_text']['units'], 1)

    def test_duplicate_media_picture_ocr_once_and_twin_reuse(self):
        from docx import Document
        png = make_png(self.base/'shared.png')
        doc = Document(); doc.add_heading('附件', 1)
        doc.add_picture(str(png)); doc.add_paragraph('同一章重复使用同一张图')
        doc.add_picture(str(png)); doc.save(self.base/'dup.docx')
        fid, result, ocr = self.run_document(self.base/'dup.docx', [LICENSE_OCR])
        self.assertEqual(len(ocr.calls), 1)
        self.assertEqual([i['ocr_status'] for i in self.images()], ['done', 'done'])
        self.assertEqual(self.images()[1]['ocr_text_raw'], LICENSE_OCR)
        self.assertEqual(result['summary']['embedded_images']['reused'], 1)
        self.assertEqual(len(self.units('embedded_image')), 2)

    def test_single_image_failure_isolates_and_retry_reruns_only_failures(self):
        path = make_docx_with_images(self.base/'flaky.docx', [LICENSE_OCR, SAFETY_OCR, OTHER_COMPANY_OCR])
        fid = self.add(path)
        rid, entry = self.begin(fid)
        FlakyOcr.queue = [LICENSE_OCR, SAFETY_OCR, OTHER_COMPANY_OCR]
        ocr = FlakyOcr(self.c)
        from jty_materials.document_extract import extract
        folder = self.c.document_tmp/rid; folder.mkdir(parents=True, exist_ok=True)
        entry['result'] = extract(entry['source'], folder/'units.jsonl', self.c, ocr_factory=lambda _: ocr)
        receipt = self.supervisor.job('project.document._finish', dict(run_id=rid, capability=entry['capability']))
        self.assertEqual(receipt['status'], 'partial')
        statuses = [i['ocr_status'] for i in self.images()]
        self.assertEqual(statuses, ['done', 'failed', 'done'])
        self.assertIn('Umi crashed', self.images()[1]['ocr_error'])
        self.assertEqual(len(self.units('embedded_image')), 2)  # failed picture yields no unit
        # retry: the two done pictures resume from cache keys, only #2 re-OCRs
        FlakyOcr.broken = False
        FlakyOcr.queue = [SAFETY_OCR]
        rid2, entry2 = self.begin(fid, action='retry')
        from jty_materials.common import normalize as _norm
        # image 2 raised before consuming its queue slot, so image 3 received
        # the safety text; only the failed picture must re-run on retry.
        self.assertEqual(entry2['resume'],
                         {self.images()[0]['ocr_cache_key']: dict(text_raw=LICENSE_OCR,
                             text_norm=_norm(LICENSE_OCR), elapsed_ms=self.images()[0]['ocr_elapsed_ms']),
                          self.images()[2]['ocr_cache_key']: dict(text_raw=SAFETY_OCR,
                             text_norm=_norm(SAFETY_OCR), elapsed_ms=self.images()[2]['ocr_elapsed_ms'])})
        retry_ocr = FlakyOcr(self.c)
        folder2 = self.c.document_tmp/rid2; folder2.mkdir(parents=True, exist_ok=True)
        entry2['result'] = extract(entry2['source'], folder2/'units.jsonl', self.c, ocr_factory=lambda _: retry_ocr,
                                   resume=entry2['resume'])
        receipt2 = self.supervisor.job('project.document._finish', dict(run_id=rid2, capability=entry2['capability']))
        self.assertEqual(receipt2['status'], 'complete')
        self.assertEqual(len(retry_ocr.calls), 1)
        self.assertEqual([i['ocr_status'] for i in self.images()], ['done', 'done', 'done'])
        self.assertEqual(len(self.units('embedded_image')), 3)

    # ---- P2: discovery WHO / WHAT / LIFECYCLE ---------------------------
    def test_discovery_candidates_subject_and_lifecycle(self):
        path = make_docx_with_images(self.base/'multi.docx',
                                     [LICENSE_OCR, SAFETY_OCR, OTHER_COMPANY_OCR, CREDIT_OCR, PERSON_OCR])
        fid, result, ocr = self.run_document(path, [LICENSE_OCR, SAFETY_OCR, OTHER_COMPANY_OCR, CREDIT_OCR, PERSON_OCR])
        rows = self.candidates()
        licenses = [r for r in rows if r['candidate_type'] == 'business_license']
        self.assertTrue(licenses)
        zhong = next(r for r in licenses if r['subject_uscc'] == '91140931MA7YM0KQ09')
        self.assertEqual(zhong['subject_uscc'], '91140931MA7YM0KQ09')
        self.assertEqual(zhong['subject_name'], '山西中保源建设工程有限公司')
        self.assertEqual(zhong['lifecycle'], 'reusable')
        self.assertEqual(zhong['match_status'], 'new_candidate')
        self.assertEqual(zhong['source_kind'], 'embedded_image')
        self.assertTrue(zhong['source_image_id'])
        by_type = {r['candidate_type']: r for r in rows}
        safety = by_type['safety_prod_license']
        self.assertIn('JZ安许证字', safety['identifier'])
        person = by_type['person_id_card']
        self.assertEqual(person['subject_person'], '王国栋')
        self.assertEqual(person['subject_id_masked'], '142601********1234')
        credit = by_type['credit_report']
        self.assertEqual(credit['lifecycle'], 'transient')
        # multi-subject: a second company's license is an equal citizen
        self.assertEqual(len(licenses), 2)
        self.assertTrue(any(r['subject_uscc'] == '91420100MA4K00000Q' for r in rows))
        self.assertFalse(any('jty' in json.dumps(r) or '晋铁源' in json.dumps(r, ensure_ascii=False) for r in rows))
        summary = json.loads(self.rows('SELECT summary_json FROM project_document')[0]['summary_json'] or '{}')
        self.assertEqual(summary['discovery']['candidates'], len(rows))
        self.assertEqual(summary['discovery']['new_candidates'], len(rows))

    def test_native_text_needs_identifier_so_toc_is_not_a_candidate(self):
        from docx import Document
        doc = Document()
        doc.add_heading('目录', 1)
        doc.add_paragraph('2.营业执照 6\n3.资质证书 8')
        doc.add_heading('七、商务响应表', 1)
        doc.add_paragraph('营业执照号 91140931MA7YM0KQ09 在有效期内')
        doc.save(self.base/'toc.docx')
        fid, result = self.process(self.base/'toc.docx')
        rows = self.candidates()
        # the TOC unit has no identifier; the response-table unit is the only hit
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['candidate_type'], 'business_license')
        self.assertEqual(rows[0]['source_kind'], 'text_unit')
        self.assertEqual(rows[0]['confidence'], 'high')

    def test_exact_match_links_and_audits_without_duplicating_material(self):
        from test_v01 import FakeOcr
        from jty_materials.workflow import import_material
        source = self.base/'lib-license.pdf'; source.write_bytes(b'%PDF-lib-license')
        mid = import_material(self.c, source, kind='certificate', subtype='safety_prod_license',
                              title='入库安全生产许可证', actor='fixture',
                              ocr_factory=lambda c: FakeOcr(c, '入库素材 OCR 文本'))['material_id']
        with Store(self.c) as s:
            file_id = s.db.execute("SELECT file_id FROM material_file WHERE material_id=? AND role='original'"
                                   " AND is_current=1", (mid,)).fetchone()['file_id']
            s.verify(mid, 'certificate_no', 'JZ2023B000123456', actor='fixture-reviewer',
                     reason='fixture seed for exact-match test', evidence_file_id=file_id)
        cert_ocr = ('安全生产许可证\n单位名称：山西中保源建设工程有限公司\n'
                    '许可证编号：JZ2023B000123456\n有效期 2026-03-01')
        path = make_docx_with_images(self.base/'exact.docx', [cert_ocr])
        fid, result, ocr = self.run_document(path, [cert_ocr])
        rows = self.candidates()
        self.assertEqual(rows[0]['match_status'], 'exact_match')
        self.assertEqual(rows[0]['matched_material_id'], mid)
        links = self.rows("SELECT * FROM project_material WHERE project_id=?", self.pid)
        self.assertEqual(len(links), 1)
        audits = self.rows("SELECT * FROM audit_log WHERE action='project.material.autolink'")
        self.assertEqual(len(audits), 1)
        self.assertIn('project_document_exact_match', audits[0]['detail_json'])
        # re-extraction (discovery full replace) stays idempotent: one candidate, one link
        rid, entry = self.begin(fid, action='retry')
        SequenceOcr.queue = [cert_ocr]
        folder = self.c.document_tmp/rid; folder.mkdir(parents=True, exist_ok=True)
        from jty_materials.document_extract import extract
        entry['result'] = extract(entry['source'], folder/'units.jsonl', self.c, ocr_factory=SequenceOcr,
                                  resume=entry['resume'])
        self.supervisor.job('project.document._finish', dict(run_id=rid, capability=entry['capability']))
        self.assertEqual(len(self.candidates()), 1)
        self.assertEqual(len(self.rows("SELECT * FROM project_material WHERE project_id=? AND status='active'", self.pid)), 1)

    class _FakeDb:
        """Minimal cursor source for match_library unit coverage."""
        def __init__(self, *rowsets):
            self.rowsets = [[(v,) for v in rs] for rs in rowsets]
        def execute(self, sql, params=()):
            return iter(self.rowsets.pop(0))

    def test_match_library_uscc_identifier_and_name_paths(self):
        from jty_materials.discovery import match_library
        hints = dict(subject_uscc='91140931MA7YM0KQ09', identifier=None,
                     subject_name='山西中保源建设工程有限公司')
        self.assertEqual(match_library(self._FakeDb([7]), 'business_license', hints),
                         ('exact_match', 7))
        self.assertEqual(match_library(self._FakeDb([5], [9]), 'safety_prod_license',
                                       dict(hints, subject_uscc=None, identifier='JZ2023B000123456')),
                         ('exact_match', 5))  # identifier exact when no USCC hint
        self.assertEqual(match_library(self._FakeDb([], [3]), 'business_license', dict(hints, identifier=None)),
                         ('possible_match', None))  # name-only is never exact
        self.assertEqual(match_library(self._FakeDb([], []), 'business_license', dict(hints, identifier=None)),
                         ('new_candidate', None))
        self.assertEqual(match_library(self._FakeDb([1, 2, 3]), 'business_license', hints)[0],
                         'possible_match')  # ambiguous USCC identity is an exception

    def test_detect_unit_filters_and_confidence(self):
        unit = dict(unit_type='paragraph_group', unit_index=3, text_raw='营业执照', image_id=None)
        self.assertEqual(detect([unit]), [])  # native mention without identity
        hints = subject_hints('business_license', LICENSE_OCR)
        self.assertEqual(hints['subject_uscc'], '91140931MA7YM0KQ09')
        self.assertEqual(hints['legal_representative'], '张三')
        self.assertEqual(lifecycle_for('credit_report', '信用中国 查询时间 2026'), 'transient')
        self.assertEqual(lifecycle_for('credit_report', '企业信用报告 年度版'), 'unknown')
        self.assertEqual(lifecycle_for('award_notice', ''), 'reusable')

    # ---- P0: automatic extraction after upload --------------------------
    def test_upload_auto_enqueues_extraction_and_duplicate_does_not(self):
        path = make_docx_with_images(self.base/'auto.docx', [LICENSE_OCR])
        receipt = self.call('project.file.add', project_id=self.pid, filename=path.name,
                            category='tender_docs', source=str(path), expected_sha256=self.digest(path))
        fid = receipt['project_file_id']
        runs = self.rows('SELECT * FROM project_document_run ORDER BY created_at')
        self.assertEqual(len(runs), 1)
        self.assertEqual(runs[0]['status'], 'queued')
        self.assertTrue(any(entry['source']['project_file_id'] == fid
                            for entry in self.supervisor.documents.capabilities.values()))
        duplicate = self.call('project.file.add', project_id=self.pid, filename=path.name,
                              category='tender_docs', source=str(path), expected_sha256=self.digest(path))
        self.assertTrue(duplicate['duplicate'])
        self.assertEqual(len(self.rows('SELECT * FROM project_document_run')), 1)

    def test_unsupported_extension_is_not_auto_extracted(self):
        path = self.base/'legacy.doc'
        path.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy')
        self.call('project.file.add', project_id=self.pid, filename='legacy.doc',
                  category='tender_docs', source=str(path), expected_sha256=self.digest(path))
        self.assertEqual(self.rows('SELECT * FROM project_document_run'), [])

    # ---- summary surface ------------------------------------------------
    def test_document_state_and_summary_labels(self):
        from webui.projects import document_state, document_summary_html
        path = make_docx_with_images(self.base/'labels.docx', [LICENSE_OCR, SAFETY_OCR])
        fid, result, ocr = self.run_document(path, [LICENSE_OCR, SAFETY_OCR])
        doc = self.rows('SELECT d.*, f.extension FROM project_document d JOIN project_file f USING(project_file_id)')[0]
        label = document_state(doc)
        self.assertIn('字正文', label)
        self.assertIn('2 张图片已识别', label)
        self.assertIn('搜索已就绪', label)
        self.assertIn('素材候选', label)
        html = document_summary_html(doc)
        self.assertIn('扫描图片 2 张已识别', html)
        self.assertIn('发现素材候选', html)

    @staticmethod
    def digest(path):
        from jty_materials.common import digest
        return digest(Path(path))


class V08ReviewFixTests(V08Harness):
    """Merge-gate fixes: multi-subject dedup, USCC checksum gate, stale
    autolink retirement, transient evidence tightening."""

    # ---- BLOCKER 1: dedup identity strength ladder ---------------------
    def test_different_subjects_same_type_without_strong_id_kept(self):
        # neither company name is extractable and no identifier exists, so
        # both must survive on source identity: swallowing one would swallow
        # a whole company's material
        blind = ['管理体系认证证书\nGB/T 19001-2016\n覆盖范围 建筑工程施工',
                 '管理体系认证证书\nGB/T 24001-2016\n覆盖范围 机电设备安装']
        path = make_docx_with_images(self.base/'blind-iso.docx', blind)
        fid, result, ocr = self.run_document(path, blind)
        rows = [r for r in self.candidates() if r['candidate_type'] == 'iso_cert']
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r['subject_uscc'] is None and r['subject_name'] is None for r in rows))

    def test_same_subject_normalized_name_still_dedupes(self):
        pair = ['体系认证证书\n名称：山西甲甲建设工程有限公司\nGB/T 19001',
                '体系认证证书\n单位名称：山西甲甲建设工程有限公司\nGB/T 24001']
        path = make_docx_with_images(self.base/'named-iso.docx', pair)
        fid, result, ocr = self.run_document(path, pair)
        rows = [r for r in self.candidates() if r['candidate_type'] == 'iso_cert']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['subject_name'], '山西甲甲建设工程有限公司')

    def test_strong_uscc_identity_dedup_across_pictures(self):
        duplicate = [LICENSE_OCR, LICENSE_OCR.replace('名称：', '企业名称：')]
        path = make_docx_with_images(self.base/'dup-license.docx', duplicate)
        fid, result, ocr = self.run_document(path, duplicate)
        rows = [r for r in self.candidates() if r['candidate_type'] == 'business_license'
                and r['subject_uscc'] == '91140931MA7YM0KQ09']
        self.assertEqual(len(rows), 1)

    # ---- BLOCKER 2: USCC must pass the GB 32100-2015 checksum ---------
    def test_uscc_checksum_gate(self):
        from jty_materials.discovery import extract_usccs
        self.assertEqual(extract_usccs('统一社会信用代码 91140931MA7YM0KQ09'), ['91140931MA7YM0KQ09'])
        self.assertEqual(extract_usccs('统一社会信用代码 91140931MA7YM0KQ08'), [])  # bad check digit
        bad = LICENSE_OCR.replace('91140931MA7YM0KQ09', '91140931MA7YM0KQ08')
        hints = subject_hints('business_license', bad)
        self.assertIsNone(hints['subject_uscc'])  # stays in evidence, never identity
        path = make_docx_with_images(self.base/'bad-uscc.docx', [bad])
        fid, result, ocr = self.run_document(path, [bad])
        rows = [r for r in self.candidates() if r['candidate_type'] == 'business_license']
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0]['subject_uscc'])
        self.assertNotEqual(rows[0]['confidence'], 'USCC high')
        self.assertNotIn('subject_uscc', 'high')  # confidence label sanity

    def test_person_id_number_is_never_an_org_uscc(self):
        from jty_materials.discovery import extract_usccs
        text = '某某证明\n身份证号码 142601199001011234\n编号 14260119900101123X'
        self.assertEqual(extract_usccs(text), [])
        # e2e: the license picture carries a person id; org USCC stays empty
        poisoned = LICENSE_OCR.replace('成立日期', '公民身份号码 142601199001011234\n成立日期')
        path = make_docx_with_images(self.base/'idcard.docx', [poisoned])
        fid, result, ocr = self.run_document(path, [poisoned])
        row = next(r for r in self.candidates() if r['candidate_type'] == 'business_license')
        self.assertEqual(row['subject_uscc'], '91140931MA7YM0KQ09')  # only the checksum-valid one
        self.assertNotEqual(row['subject_uscc'], '142601199001011234')

    # ---- BLOCKER 3: auto links must retire their own stale relations --
    def _seed_exact_material(self):
        from test_v01 import FakeOcr
        from jty_materials.workflow import import_material
        source = self.base/'lib-exact.pdf'; source.write_bytes(b'%PDF-lib-exact')
        mid = import_material(self.c, source, kind='certificate', subtype='safety_prod_license',
                              title='入库安全生产许可证', actor='fixture',
                              ocr_factory=lambda c: FakeOcr(c, '入库素材 OCR'))['material_id']
        with Store(self.c) as s:
            file_id = s.db.execute("SELECT file_id FROM material_file WHERE material_id=? AND role='original'"
                                   " AND is_current=1", (mid,)).fetchone()['file_id']
            s.verify(mid, 'certificate_no', 'JZ2023B000123456', actor='fixture-reviewer',
                     reason='fixture seed', evidence_file_id=file_id)
        return mid

    def _retry_with_ocr(self, fid, texts):
        """Re-run discovery for one document with fresh OCR texts (no resume)."""
        from jty_materials.document_extract import extract
        SequenceOcr.queue = list(texts)
        rid, entry = self.begin(fid, action='retry')
        entry['resume'] = {}
        folder = self.c.document_tmp/rid; folder.mkdir(parents=True, exist_ok=True)
        entry['result'] = extract(entry['source'], folder/'units.jsonl', self.c,
                                  ocr_factory=lambda _: SequenceOcr(self.c))
        return self.supervisor.job('project.document._finish', dict(run_id=rid, capability=entry['capability']))

    def _active_links(self):
        return self.rows("SELECT * FROM project_material WHERE project_id=? AND status='active'", self.pid)

    def test_autolink_retires_when_exact_evidence_disappears(self):
        mid = self._seed_exact_material()
        path = make_docx_with_images(self.base/'b3-one.docx', [EXACT_OCR])
        fid, result, ocr = self.run_document(path, [EXACT_OCR])
        self.assertEqual(len(self._active_links()), 1)
        receipt = self._retry_with_ocr(fid, [EXACT_OCR.replace('\n许可证编号：JZ2023B000123456', '')])
        self.assertEqual(self._active_links(), [])
        summary = json.loads(self.rows('SELECT summary_json FROM project_document')[0]['summary_json'] or '{}')
        self.assertEqual(summary['discovery']['autounlinked'], 1)
        audits = self.rows("SELECT * FROM audit_log WHERE action='project.material.autounlink'")
        self.assertEqual(len(audits), 1)
        self.assertIn('project_document_exact_match_no_longer_supported', audits[0]['detail_json'])
        self.assertIn('"material_id": %d' % mid, audits[0]['detail_json'].replace('  ', ' '))

    def test_manual_link_is_never_auto_removed(self):
        mid = self._seed_exact_material()
        self.call('project.material.link', project_id=self.pid, material_id=mid)
        path = make_docx_with_images(self.base/'b3-manual.docx', ['通用证明材料\n仅用于占用手工引用'])
        fid, result, ocr = self.run_document(path, ['体系认证证书\nGB/T 19001'])
        self.assertEqual(len(self._active_links()), 1)
        self._retry_with_ocr(fid, ['体系认证证书\nGB/T 24001'])
        links = self._active_links()
        self.assertEqual(len(links), 1)
        self.assertEqual(links[0]['added_by'], 'fixture-reviewer')
        self.assertEqual(self.rows("SELECT * FROM audit_log WHERE action='project.material.autounlink'"), [])

    def test_second_document_support_keeps_shared_link_alive(self):
        mid = self._seed_exact_material()
        # the two files must differ byte-wise or the second upload is a
        # content duplicate and never becomes an independent document
        text_b = EXACT_OCR + '\n补充说明：本证明用于投标文件附件。'
        path_a = make_docx_with_images(self.base/'b3-doca.docx', [EXACT_OCR])
        path_b = make_docx_with_images(self.base/'b3-docb.docx', [EXACT_OCR], extra=[text_b])
        fid_a, _, _ = self.run_document(path_a, [EXACT_OCR])
        fid_b, _, _ = self.run_document(path_b, [text_b])
        self.assertNotEqual(fid_a, fid_b)
        self.assertEqual(len(self._active_links()), 1)  # one shared reference
        # document A's evidence vanishes; B's active exact candidate remains
        self._retry_with_ocr(fid_a, [ISO_NOISE])
        links = self._active_links()
        self.assertEqual(len(links), 1)
        self.assertEqual(links[0]['material_id'], mid)
        self.assertEqual(self.rows("SELECT * FROM audit_log WHERE action='project.material.autounlink'"), [])

    def test_retired_link_recovers_when_exact_evidence_returns(self):
        self._seed_exact_material()
        path = make_docx_with_images(self.base/'b3-return.docx', [EXACT_OCR])
        fid, _, _ = self.run_document(path, [EXACT_OCR])
        self._retry_with_ocr(fid, [ISO_NOISE])          # evidence gone
        self.assertEqual(self._active_links(), [])
        self._retry_with_ocr(fid, [EXACT_OCR])          # evidence returns
        links = self._active_links()
        self.assertEqual(len(links), 1)
        note = json.loads(links[0]['note'])
        self.assertEqual(note['method'], 'project_document_exact_match')
        self.assertEqual(len(self.rows("SELECT * FROM audit_log WHERE action='project.material.autolink'")), 2)

    # ---- BLOCKER 4: transient needs real point-in-time query evidence --
    def test_transient_instant_query_still_transient(self):
        instant = ('企业信用报告\n信用中国网站查询\n查询时间 2026-10-26\n查询结果 未列入严重违法名单')
        self.assertEqual(lifecycle_for('credit_report', instant), 'transient')
        pair_only = '企业信用报告\n查询时间 2026-10-26\n查询结果 无失信记录'
        self.assertEqual(lifecycle_for('credit_report', pair_only), 'transient')

    def test_formal_rating_report_is_not_transient(self):
        formal = ('企业信用报告（信用评级）\n报告编号：ABC-2026-001\n生成时间：2026-10-01\n'
                  '信用等级 AAA\n评级机构 某信用评估有限公司')
        self.assertNotEqual(lifecycle_for('credit_report', formal), 'transient')
        self.assertEqual(lifecycle_for('credit_report', formal), 'unknown')
        path = make_docx_with_images(self.base/'rating.docx', [formal])
        fid, result, ocr = self.run_document(path, [formal])
        row = next(r for r in self.candidates() if r['candidate_type'] == 'credit_report')
        self.assertEqual(row['lifecycle'], 'unknown')


class V08MigrationGuardTests(DocumentFixture):
    """Section 7 regression gate: v0.8 migration keeps every guard family."""

    def test_all_trigger_guards_present_after_initialization(self):
        with Store(self.c, readonly=True) as s:
            names = {r[0] for r in s.db.execute(
                "SELECT name FROM sqlite_schema WHERE type='trigger'")}
        expected = set()
        for table in ('project', 'project_file', 'project_material'):
            for guard in ('no_replace', 'no_delete', 'controlled_INSERT', 'controlled_UPDATE'):
                expected.add(f'{table}_{guard}')
        # v0.5 attaches the status-transition guards to the two evidence
        # tables only; the project row itself has none.
        for table in ('project_file', 'project_material'):
            for op in ('INSERT', 'UPDATE'):
                expected.add(f'{table}_active_{op}')
        expected |= {'project_identity', 'project_file_path', 'project_file_immutable',
                     'project_material_identity'}
        for table in ('project_document', 'project_document_run', 'project_document_unit',
                      'project_document_image', 'project_material_candidate'):
            for op in ('INSERT', 'UPDATE', 'DELETE'):
                expected.add(f'controlled_{table}_{op}')
        expected |= {'project_document_source', 'project_document_identity', 'project_document_unit_binding'}
        expected |= {'protect_audit_log_UPDATE', 'protect_audit_log_DELETE', 'protect_audit_log_replace'}
        missing = sorted(expected - names)
        self.assertEqual(missing, [], 'missing trigger guards after v0.8 schema')

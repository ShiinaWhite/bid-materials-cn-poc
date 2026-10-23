"""Small deterministic generators; all generated binaries live under runs."""
import io
import json
import os
from pathlib import Path
import tempfile
from unittest import mock

import pymupdf
from PIL import Image, ImageDraw, ImageFont

from project_test_support import ProjectFixture
from jty_materials.common import digest
from jty_materials.document_extract import extract


def make_pdf(path, kinds=('native',), term='PROJECT_DOC_SENTINEL_654321'):
    with pymupdf.open() as doc:
        for i, kind in enumerate(kinds):
            page = doc.new_page(width=595, height=842)
            if kind == 'native':
                page.insert_text((40, 60), f'{term} page {i+1} project manager payment terms quality standard', fontsize=10)
                page.insert_text((40, 85), '项目经理 工期 质量标准 付款方式 安全生产许可证 10kV', fontname='china-s', fontsize=12)
            elif kind == 'scan':
                image = Image.new('RGB', (1400, 1900), 'white')
                draw = ImageDraw.Draw(image)
                font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 48)
                draw.text((80, 150), 'PROJECT MANAGER PAYMENT', font=font, fill='black')
                draw.text((80, 250), 'SAFETY CERTIFICATE 998877', font=font, fill='black')
                chinese = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 52)
                draw.text((80, 370), '项目经理 工期 质量标准', font=chinese, fill='black')
                output = io.BytesIO(); image.save(output, format='PNG')
                page.insert_image(page.rect, stream=output.getvalue())
        doc.save(path, deflate=True)
    return path


def make_docx(path, term='DOCX_SECTION_SENTINEL', sections=1):
    from docx import Document
    doc = Document()
    for i in range(sections):
        doc.add_heading(f'第四章 投标文件格式 {i}', 1)
        doc.add_heading('五、项目管理机构', 2)
        doc.add_heading('（一）项目经理', 3)
        doc.add_paragraph(term + ' 项目经理须具备有效资格，仅为历史参考。')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0,0).text = '付款方式'
        table.cell(0,1).text = '按月支付'
        table.cell(1,0).text = '质量标准'
        table.cell(1,1).text = '合格'
    doc.save(path)
    return path


def make_xlsx(path, rows=80, sheets=2):
    import openpyxl
    doc = openpyxl.Workbook()
    doc.remove(doc.active)
    for i in range(sheets):
        sheet = doc.create_sheet('分部分项工程量清单' + str(i+1))
        sheet.append(['名称','数量','备注'])
        for r in range(2, rows+1):
            if r not in (20, 21, 22):
                sheet.cell(r,1,'10kV 电缆 XLSX_SHEET_SENTINEL')
                sheet.cell(r,2,r)
                sheet.cell(r,3,f'=B{r}*2')
        sheet.cell(rows+200,1,'SPARSE_ROW_SENTINEL')
    doc.save(path)
    return path


class StubOcr:
    def __init__(self, config): self.calls = 0
    def recognize(self, data):
        self.calls += 1
        return 'SCANNED_OCR_SENTINEL 项目经理'
    def close(self): pass


class DocumentFixture(ProjectFixture):
    def setUp(self):
        temp = Path(__file__).resolve().parents[1]/'runs/large-document/tmp'
        temp.mkdir(parents=True, exist_ok=True)
        os.environ['TEMP'] = os.environ['TMP'] = str(temp)
        tempfile.tempdir = str(temp)
        super().setUp()
        self.pid = self.create()
        self.enqueue = mock.patch.object(self.supervisor.documents, 'enqueue')
        self.queued = self.enqueue.start()
        self.addCleanup(self.enqueue.stop)

    def add(self, path, pid=None):
        return self.call('project.file.add', project_id=pid or self.pid, filename=path.name,
            category='tender_docs', source=str(path), expected_sha256=digest(path))['project_file_id']

    def begin(self, fid, pid=None, action='extract'):
        result = self.call('project.document.'+action, project_id=pid or self.pid, project_file_id=fid)
        rid = result['run_id']
        return rid, self.supervisor.documents.capabilities[rid]

    def complete(self, rid, entry, ocr_factory=StubOcr):
        folder = self.c.document_tmp/rid
        folder.mkdir(parents=True, exist_ok=True)
        entry['result'] = extract(entry['source'], folder/'units.jsonl', self.c, ocr_factory=ocr_factory)
        return self.supervisor.job('project.document._finish', dict(run_id=rid, capability=entry['capability']))

    def process(self, path, pid=None):
        fid = self.add(path, pid)
        rid, entry = self.begin(fid, pid)
        result = self.complete(rid, entry)
        return fid, result
